"""PDF -> Excel: cell typing, table heuristics, and end-to-end conversions of generated PDFs."""

import tempfile
import unittest
from pathlib import Path

import pymupdf
from openpyxl import load_workbook

from app.tables import Table, clean_rows, detect_decimal, merge_continued, merge_unit_columns, plausible, type_cell


class TypeCell(unittest.TestCase):
    def check(self, raw, decimal, value, fmt=None):
        t = type_cell(raw, decimal)
        self.assertIsNotNone(t, raw)
        self.assertAlmostEqual(t.value, value, msg=raw)
        if fmt:
            self.assertEqual(t.number_format, fmt, raw)

    def test_french_numbers(self):
        self.check("1 234,56", ",", 1234.56, "#,##0.00")
        self.check("1 234,5", ",", 1234.5)
        self.check("1.234.567,8", ",", 1234567.8)
        self.check("3,50 €", ",", 3.5, '#,##0.00 "€"')
        self.check("-45,00 €", ",", -45.0)
        self.check("(1 200)", ",", -1200)
        self.check("12 %", ",", 0.12, "0%")
        self.check("12,5 %", ",", 0.125, "0.0%")

    def test_english_numbers(self):
        self.check("1,234.56", ".", 1234.56)
        self.check("$1,200.00", ".", 1200.0, '"$"#,##0.00')
        self.check("3.14", ".", 3.14)
        self.check("42", ".", 42, "General")
        self.check("−7", ".", -7)  # unicode minus

    def test_kept_as_text(self):
        for raw in ["0612345678", "06 12 34 56 78", "00123", "12/03/2024", "abc", "FR76 3000 6000",
                    "4111111111111111111", "1.2.3", "", "N° 12a"]:
            self.assertIsNone(type_cell(raw, ","), raw)


class Heuristics(unittest.TestCase):
    def test_detect_decimal(self):
        self.assertEqual(detect_decimal(["1 234,56", "3,5", "Total"]), ",")
        self.assertEqual(detect_decimal(["1,234.56", "3.25"]), ".")
        self.assertEqual(detect_decimal(["1,234", "abc"], default=","), ",")  # ambiguous -> default

    def test_clean_rows_drops_empty(self):
        self.assertEqual(clean_rows([["a", None, "b"], [None, None, None], ["c", "", " d \n e"]]), [["a", "b"], ["c", "d e"]])

    def test_unit_columns_glued_back(self):
        rows = [["Prix", "", "Remise"], ["3,50", "€", "10 %"], ["4,20", "€", "5 %"]]
        self.assertEqual(merge_unit_columns(rows), [["Prix", "Remise"], ["3,50 €", "10 %"], ["4,20 €", "5 %"]])

    def test_plausible(self):
        self.assertTrue(plausible([["Date", "Montant"], ["01/02", "12,00"], ["03/02", "8,50"]]))
        paragraph = [["Ceci est une longue phrase qui n'est clairement pas un tableau de données", ""],
                     ["et elle continue encore sur une deuxième ligne de texte assez longue", ""]]
        self.assertFalse(plausible(paragraph))

    def test_merge_continued_drops_repeated_header(self):
        a = Table([["H1", "H2"], ["1", "2"]], 1, 1, True, True)
        b = Table([["H1", "H2"], ["3", "4"]], 2, 2, True, True)
        c = Table([["X", "Y", "Z"], ["5", "6", "7"]], 3, 3, True, True)  # other columns: separate table
        merged = merge_continued([a, b, c])
        self.assertEqual(len(merged), 2)
        self.assertEqual(merged[0].rows, [["H1", "H2"], ["1", "2"], ["3", "4"]])
        self.assertEqual((merged[0].page, merged[0].last_page), (1, 2))


def _draw_table(page, x, y, rows, widths, ruled, row_h=22):
    for i, row in enumerate(rows):
        cx = x
        for text, w in zip(row, widths):
            page.insert_text((cx + 4, y + i * row_h + 15), text, fontsize=10, fontname="dejavu", fontfile=FONT)
            cx += w
    if ruled:
        total_w = sum(widths)
        for i in range(len(rows) + 1):
            page.draw_line((x, y + i * row_h), (x + total_w, y + i * row_h))
        cx = x
        for w in [0] + widths:
            cx += w
            page.draw_line((cx, y), (cx, y + len(rows) * row_h))


FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"  # base-14 Helvetica can't encode €
HEADER = ["Produit", "Quantité", "Prix", "Remise"]
DATA = [["Pommes", "1 200", "3,50 €", "10 %"], ["Poires", "850", "4,20 €", "5 %"], ["Kiwis", "2 300", "0,90 €", "0 %"]]
WIDTHS = [120, 90, 90, 80]


class EndToEnd(unittest.TestCase):
    def convert(self, doc: pymupdf.Document):
        from app.tables import convert

        with tempfile.TemporaryDirectory() as tmp:
            src, out = Path(tmp) / "in.pdf", Path(tmp) / "out.xlsx"
            doc.save(src)
            convert(str(src), str(out), tmp)
            wb = load_workbook(out)
            return {ws.title: [[c.value for c in r] for r in ws.iter_rows()] for ws in wb.worksheets}

    def test_ruled_table_typed(self):
        doc = pymupdf.open()
        _draw_table(doc.new_page(), 50, 80, [HEADER] + DATA, WIDTHS, ruled=True)
        sheets = self.convert(doc)
        rows = sheets["Table 1 (p1)"]
        self.assertEqual(rows[0], HEADER)
        self.assertEqual(rows[1], ["Pommes", 1200, 3.5, 0.1])

    def test_borderless_table_and_paragraph_ignored(self):
        doc = pymupdf.open()
        page = doc.new_page()
        page.insert_textbox(pymupdf.Rect(50, 40, 545, 120),
                            "Ce rapport présente les ventes du trimestre. Les chiffres ci-dessous sont indicatifs "
                            "et peuvent évoluer après la clôture des comptes.", fontsize=10)
        _draw_table(page, 50, 150, [HEADER] + DATA, WIDTHS, ruled=False)
        rows = next(iter(self.convert(doc).values()))
        self.assertIn(HEADER, rows)
        self.assertIn(["Kiwis", 2300, 0.9, 0], rows)
        self.assertFalse(any("Ce rapport" in str(c) for r in rows for c in r if c))

    def test_table_continued_on_next_page(self):
        doc = pymupdf.open()
        _draw_table(doc.new_page(), 50, 80, [HEADER] + DATA, WIDTHS, ruled=True)
        _draw_table(doc.new_page(), 50, 80, [HEADER, ["Prunes", "40", "6,00 €", "15 %"]], WIDTHS, ruled=True)
        sheets = self.convert(doc)
        self.assertEqual(list(sheets), ["Table 1 (p1-2)"])
        rows = sheets["Table 1 (p1-2)"]
        self.assertEqual(len(rows), 5)  # header once + 4 data rows
        self.assertEqual(rows[-1], ["Prunes", 40, 6.0, 0.15])

    def test_scanned_pdf_is_ocrd(self):
        ruled = pymupdf.open()
        _draw_table(ruled.new_page(), 50, 80, [HEADER] + DATA, WIDTHS, ruled=True)
        pix = ruled[0].get_pixmap(dpi=200)
        scanned = pymupdf.open()
        page = scanned.new_page(width=ruled[0].rect.width, height=ruled[0].rect.height)
        page.insert_image(page.rect, pixmap=pix)  # image only: no text layer
        rows = next(iter(self.convert(scanned).values()))
        flat = [str(c) for r in rows for c in r if c is not None]
        self.assertTrue(any("Pommes" in c for c in flat), flat)
        self.assertIn(1200, [c for r in rows for c in r])


if __name__ == "__main__":
    unittest.main()
