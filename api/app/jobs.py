"""CPU-heavy jobs, executed in isolated processes by `isolate.run_job`.

Every function here is module-level and takes plain values (paths as str), so it can be pickled
by reference. Importing this module also preloads the heavy libraries in the forkserver.
"""

import difflib
import os
from pathlib import Path

import pikepdf
import pymupdf

from .errors import ApiError
from .redact import redact  # noqa: F401  (re-exported as a job)
from .util import assert_public_url

MAX_PAGES = int(os.getenv("MAX_PAGES", "500"))
UPSCALE_MAX_INPUT_PIXELS = {2: 2_500_000, 4: 800_000}


def check_pages(src: str, limit: int = MAX_PAGES) -> int:
    """Reject PDFs with more pages than `limit` before handing them to a slow tool."""
    with pymupdf.open(src) as doc:
        if doc.needs_pass:
            raise ApiError(400, "encrypted", "PDF is password protected")
        if doc.page_count == 0:
            raise ApiError(400, "no_pages", "PDF has no pages")
        if doc.page_count > limit:
            raise ApiError(413, "too_many_pages", "Too many pages", max=limit)
        return doc.page_count


def repair(src: str, out: str) -> bool:
    """Rewrite through qpdf; returns False when the file is too broken (caller falls back to Ghostscript)."""
    try:
        with pikepdf.open(src) as pdf:
            pdf.save(out)
        return True
    except pikepdf.PasswordError:
        raise
    except Exception:
        return False


def pdf_to_word(src: str, out: str) -> None:
    from pdf2docx import Converter

    cv = Converter(src)
    try:
        cv.convert(out)
    finally:
        cv.close()


def pdf_to_excel(src: str, out: str) -> None:
    import pdfplumber
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    with pdfplumber.open(src) as pdf:
        for pno, page in enumerate(pdf.pages, start=1):
            tables = page.extract_tables()
            if tables:
                for tno, table in enumerate(tables, start=1):
                    sheet = wb.create_sheet(f"P{pno}-T{tno}"[:31])
                    for row in table:
                        sheet.append([c if c is not None else "" for c in row])
            else:
                text = page.extract_text() or ""
                if text.strip():
                    sheet = wb.create_sheet(f"P{pno}-text"[:31])
                    for line in text.splitlines():
                        sheet.append(line.split())
    if not wb.sheetnames:
        wb.create_sheet("Empty")
    wb.save(out)


def pdf_to_powerpoint(src: str, out: str, tmp_dir: str) -> None:
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation()
    with pymupdf.open(src) as doc:
        first = doc[0].rect
        # 1pt = 12700 EMU
        prs.slide_width, prs.slide_height = Emu(int(first.width * 12700)), Emu(int(first.height * 12700))
        blank = prs.slide_layouts[6]
        for i, page in enumerate(doc):
            img = Path(tmp_dir) / f"slide{i}.png"
            page.get_pixmap(dpi=150).save(img)
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(str(img), 0, 0, prs.slide_width, prs.slide_height)
    prs.save(out)


def _safe_fetcher():
    """WeasyPrint fetcher limited to data: and public http(s) URLs, re-checking every redirect hop."""
    from urllib import request

    from weasyprint.urls import URLFetcher

    def check(url: str) -> None:
        try:
            assert_public_url(url)
        except ApiError as e:
            raise ValueError(e.detail["message"])

    class _CheckedRedirects(request.HTTPRedirectHandler):
        def redirect_request(self, req, fp, code, msg, headers, newurl):
            check(newurl)
            return super().redirect_request(req, fp, code, msg, headers, newurl)

    class SafeFetcher(URLFetcher):
        def __init__(self):
            super().__init__(timeout=15, allowed_protocols={"http", "https", "data"}, allow_redirects=False)
            self.add_handler(_CheckedRedirects())

        def fetch(self, url, headers=None):
            if not url.startswith("data:"):
                check(url)
            return super().fetch(url, headers)

    return SafeFetcher()


def html_to_pdf(out: str, url: str | None = None, html_path: str | None = None) -> None:
    from weasyprint import HTML

    if url:
        HTML(url=url, url_fetcher=_safe_fetcher()).write_pdf(out)
    else:
        html = Path(html_path).read_text(errors="replace")
        HTML(string=html, url_fetcher=_safe_fetcher()).write_pdf(out)


def protect(src: str, out: str, password: str) -> None:
    with pikepdf.open(src) as pdf:
        pdf.save(out, encryption=pikepdf.Encryption(owner=password, user=password, R=6))


def unlock(src: str, out: str, password: str) -> None:
    try:
        with pikepdf.open(src, password=password) as pdf:
            pdf.save(out)
    except pikepdf.PasswordError:
        raise ApiError(400, "wrong_password", "Wrong password")


def flatten(src: str, out: str) -> None:
    with pikepdf.open(src) as pdf:
        # form fields need appearance streams before they can be burnt into the page
        pdf.generate_appearance_streams()
        pdf.flatten_annotations("all")
        if "/AcroForm" in pdf.Root:
            del pdf.Root.AcroForm
        pdf.save(out)


def _words(doc: pymupdf.Document) -> list[tuple[int, pymupdf.Rect, str]]:
    return [(pno, pymupdf.Rect(w[:4]), w[4]) for pno, page in enumerate(doc) for w in page.get_text("words")]


def compare(src_a: str, src_b: str, out: str, name_a: str, name_b: str) -> None:
    with pymupdf.open(src_a) as a, pymupdf.open(src_b) as b, pymupdf.open() as result:
        wa, wb = _words(a), _words(b)
        removed: dict[int, list[pymupdf.Rect]] = {}
        added: dict[int, list[pymupdf.Rect]] = {}
        sm = difflib.SequenceMatcher(None, [w[2] for w in wa], [w[2] for w in wb], autojunk=False)
        for tag, i1, i2, j1, j2 in sm.get_opcodes():
            if tag in ("delete", "replace"):
                for pno, rect, _ in wa[i1:i2]:
                    removed.setdefault(pno, []).append(rect)
            if tag in ("insert", "replace"):
                for pno, rect, _ in wb[j1:j2]:
                    added.setdefault(pno, []).append(rect)

        gap, header = 24, 28
        for i in range(max(a.page_count, b.page_count)):
            ra = a[i].rect if i < a.page_count else pymupdf.Rect(0, 0, 0, 0)
            rb = b[i].rect if i < b.page_count else pymupdf.Rect(0, 0, 0, 0)
            width = ra.width + rb.width + gap
            height = max(ra.height, rb.height) + header
            page = result.new_page(width=width, height=height)
            page.insert_text((8, 18), f"A: {name_a}  (page {i + 1})", fontsize=10, color=(0.7, 0, 0))
            page.insert_text((ra.width + gap + 8, 18), f"B: {name_b}  (page {i + 1})", fontsize=10, color=(0, 0.5, 0))
            if i < a.page_count:
                page.show_pdf_page(pymupdf.Rect(0, header, ra.width, header + ra.height), a, i)
                for r in removed.get(i, []):
                    page.draw_rect(r + (0, header, 0, header), color=None, fill=(1, 0, 0), fill_opacity=0.3)
            if i < b.page_count:
                ox = ra.width + gap
                page.show_pdf_page(pymupdf.Rect(ox, header, ox + rb.width, header + rb.height), b, i)
                for r in added.get(i, []):
                    page.draw_rect(r + (ox, header, ox, header), color=None, fill=(0, 0.8, 0), fill_opacity=0.3)
            page.draw_line((ra.width + gap / 2, 0), (ra.width + gap / 2, height), color=(0.6, 0.6, 0.6))
        result.save(out)


def upscale(src: str, out: str, scale: int) -> None:
    from PIL import Image, ImageFilter

    with Image.open(src) as im:
        if im.width * im.height > UPSCALE_MAX_INPUT_PIXELS[scale]:
            raise ApiError(400, "image_too_large", "Image too large", scale=scale)
        im = im.convert("RGBA") if im.mode in ("RGBA", "LA", "P") else im.convert("RGB")
        model = Path(os.getenv("MODELS_DIR", "/models")) / f"FSRCNN_x{scale}.pb"
        try:
            import cv2
            import numpy as np

            if not model.exists():
                raise FileNotFoundError(model)
            sr = cv2.dnn_superres.DnnSuperResImpl_create()
            sr.readModel(str(model))
            sr.setModel("fsrcnn", scale)
            rgb = np.array(im.convert("RGB"))[:, :, ::-1]
            up = Image.fromarray(sr.upsample(rgb)[:, :, ::-1])
            if im.mode == "RGBA":
                up.putalpha(im.getchannel("A").resize(up.size, Image.LANCZOS))
        except (ImportError, AttributeError, FileNotFoundError):
            up = im.resize((im.width * scale, im.height * scale), Image.LANCZOS)
            up = up.filter(ImageFilter.UnsharpMask(radius=2, percent=80, threshold=2))
        up.save(out, optimize=True)
