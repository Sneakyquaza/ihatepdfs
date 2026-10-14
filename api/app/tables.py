"""PDF tables -> Excel workbook.

Pipeline: extract tables per page (ruled lines first, then text alignment with a plausibility check)
-> merge tables continued across pages -> type the cells (numbers, percentages, currencies) -> write
a workbook with one sheet per table. Scanned PDFs (no text layer) are OCR'd first. If no table is
found at all, the text is laid out in columns as a fallback sheet.
"""

import re
import subprocess
from dataclasses import dataclass
from pathlib import Path

# ---------------------------------------------------------------------------
# Cell typing (pure, unit-tested)

_SPACES = "     "
_MINUS = "-−–"
_CURRENCIES = {"€": "€", "$": "$", "£": "£", "EUR": "€", "USD": "$", "GBP": "£", "CHF": "CHF"}
_CURRENCY_RE = re.compile(r"^(€|\$|£|EUR|USD|GBP|CHF)\s*|\s*(€|\$|£|EUR|USD|GBP|CHF)$")
_NUMERIC_SHAPE = re.compile(r"^[(" + _MINUS + r"+]?\s*(?:€|\$|£)?\s*[\d" + _SPACES + r".,']+\s*(?:€|\$|£|%|EUR|USD|GBP|CHF)?\s*[)" + _MINUS + r"]?$")
MAX_DIGITS = 15  # Excel keeps 15 significant digits; longer runs are identifiers


@dataclass(frozen=True)
class Typed:
    value: float | int
    number_format: str


def detect_decimal(cells: list[str], default: str = ",") -> str:
    """Decimal separator used by a table, from its unambiguous numbers ("1,5", "1.234,56", "3.14")."""
    votes = {",": 0, ".": 0}
    for raw in cells:
        s = raw.strip().strip("".join(_SPACES))
        if not s or not _NUMERIC_SHAPE.match(s):
            continue
        digits = re.sub(r"[^\d.,]", "", s)
        if "," in digits and "." in digits:
            votes[digits[max(digits.rfind(","), digits.rfind("."))]] += 1
            continue
        for sep in ",.":
            m = re.search(r"\d\{}(\d+)$".format(sep), digits)
            if m and digits.count(sep) == 1 and len(m.group(1)) != 3:  # x,ddd / x.ddd stay ambiguous
                votes[sep] += 1
    if votes[","] == votes["."]:
        return default
    return "," if votes[","] > votes["."] else "."


def type_cell(raw: str, decimal: str) -> Typed | None:
    """Number/percent/currency value of a cell, or None to keep it as text."""
    s = raw.strip()
    if not s or not _NUMERIC_SHAPE.match(s):
        return None
    negative = False
    if s.startswith("(") and s.endswith(")"):
        negative, s = True, s[1:-1].strip()
    if s and s[0] in _MINUS:
        negative, s = True, s[1:].strip()
    elif s and s[-1] in _MINUS:
        negative, s = True, s[:-1].strip()
    elif s.startswith("+"):
        s = s[1:].strip()
    percent = s.endswith("%")
    if percent:
        s = s[:-1].strip()
    currency, prefix = None, False
    m = _CURRENCY_RE.search(s)
    if m:
        currency = _CURRENCIES[m.group(1) or m.group(2)]
        prefix = bool(m.group(1))
        s = _CURRENCY_RE.sub("", s).strip()
    if s and s[0] in _MINUS:  # "€ -12" / "-12 €"
        negative, s = True, s[1:].strip()

    thousands = "." if decimal == "," else ","
    compact = s.translate({ord(c): None for c in _SPACES + "'"})
    if not re.fullmatch(r"\d[\d" + re.escape(thousands) + r"]*(?:" + re.escape(decimal) + r"\d+)?", compact):
        return None
    int_part, _, frac = compact.partition(decimal)
    groups = int_part.split(thousands)
    if len(groups) > 1 and (len(groups[0]) > 3 or any(len(g) != 3 for g in groups[1:])):
        return None  # "12.34.5" is not a thousands-grouped number
    int_digits = "".join(groups)
    if len(int_digits) > 1 and int_digits.startswith("0") and not frac:
        return None  # 00123, 0612345678: identifiers, keep the leading zeros
    if len(int_digits) + len(frac) > MAX_DIGITS:
        return None

    value: float | int = float(f"{int_digits}.{frac}") if frac else int(int_digits)
    if negative:
        value = -value
    grouped = len(groups) > 1 or any(c in s for c in _SPACES + "'")
    decimals = "0" * len(frac)
    base = ("#,##0" if grouped or currency else "0") + (f".{decimals}" if decimals else "")
    if percent:
        return Typed(value / 100, f"0{'.' + decimals if decimals else ''}%")
    if currency:
        sym = f'"{currency}"'
        return Typed(value, f"{sym}{base}" if prefix else f"{base} {sym}")
    return Typed(value, base if (grouped or decimals) else "General")


# ---------------------------------------------------------------------------
# Table structure (pure, unit-tested)

@dataclass
class Table:
    rows: list[list[str]]
    page: int  # 1-based
    last_page: int
    first_on_page: bool
    last_on_page: bool

    @property
    def ncols(self) -> int:
        return max((len(r) for r in self.rows), default=0)


def clean_rows(raw: list[list[str | None]]) -> list[list[str]]:
    """Normalize cells, drop empty rows and columns."""
    rows = [[" ".join((c or "").split()) for c in r] for r in raw]
    rows = [r for r in rows if any(r)]
    if not rows:
        return []
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    keep = [j for j in range(width) if any(r[j] for r in rows)]
    return [[r[j] for j in keep] for r in rows]


_UNIT = re.compile(r"^(?:€|\$|£|%|EUR|USD|GBP|CHF|HT|TTC)$", re.IGNORECASE)


def merge_unit_columns(rows: list[list[str]]) -> list[list[str]]:
    """Text-aligned extraction splits "3,50 €" into "3,50" | "€": glue unit-only columns back to the left."""
    if not rows:
        return rows
    ncols = len(rows[0])
    unit_cols = [j for j in range(1, ncols)
                 if any(r[j] for r in rows) and all(not r[j] or _UNIT.match(r[j]) for r in rows)]
    if not unit_cols:
        return rows
    out = []
    for r in rows:
        row = list(r)
        for j in reversed(unit_cols):
            if row[j]:
                row[j - 1] = f"{row[j - 1]} {row[j]}".strip()
            del row[j]
        out.append(row)
    return out


def plausible(rows: list[list[str]]) -> bool:
    """Reject 'tables' that are really paragraphs split on word gaps (text-alignment strategy)."""
    if len(rows) < 2 or not rows or len(rows[0]) < 2:
        return False
    cells = [c for r in rows for c in r]
    filled = [c for c in cells if c]
    if len(filled) / len(cells) < 0.5:
        return False
    if sum(len(c) for c in filled) / len(filled) > 40:
        return False
    # at least two columns that are mostly filled
    ncols = len(rows[0])
    dense = sum(1 for j in range(ncols) if sum(1 for r in rows if r[j]) >= len(rows) / 2)
    return dense >= 2


def merge_continued(tables: list[Table]) -> list[Table]:
    """A table ending a page and one starting the next page with the same columns are one table."""
    merged: list[Table] = []
    for t in tables:
        prev = merged[-1] if merged else None
        if prev and prev.last_on_page and t.first_on_page and t.page == prev.last_page + 1 and t.ncols == prev.ncols:
            rows = t.rows[1:] if t.rows and prev.rows and t.rows[0] == prev.rows[0] else t.rows  # repeated header
            prev.rows.extend(rows)
            prev.last_page = t.page
            prev.last_on_page = t.last_on_page
        else:
            merged.append(t)
    return merged


def is_header(row: list[str], decimal: str) -> bool:
    filled = [c for c in row if c]
    return len(filled) >= max(2, len(row) // 2) and all(type_cell(c, decimal) is None for c in filled)


# ---------------------------------------------------------------------------
# Extraction and workbook

TEXT_STRATEGY = {
    "vertical_strategy": "text",
    "horizontal_strategy": "text",
    "snap_tolerance": 3,
    "join_tolerance": 3,
    "min_words_vertical": 3,
    "min_words_horizontal": 2,
    "text_x_tolerance": 2,
}


def _page_tables(page, pno: int) -> list[Table]:
    found = [merge_unit_columns(clean_rows(t.extract())) for t in page.find_tables()]
    found = [rows for rows in found if len(rows) >= 2 and len(rows[0]) >= 2]
    if not found and page.chars:
        candidates = [merge_unit_columns(clean_rows(t.extract())) for t in page.find_tables(TEXT_STRATEGY)]
        found = [rows for rows in candidates if plausible(rows)]
    return [Table(rows, pno, pno, i == 0, i == len(found) - 1) for i, rows in enumerate(found)]


def _ocr_if_scanned(src: str, tmp_dir: str) -> str:
    """Scanned PDFs have no text layer: add one with OCR so tables can be read."""
    import pdfplumber

    with pdfplumber.open(src) as pdf:
        if any(page.chars for page in pdf.pages):
            return src
    out = str(Path(tmp_dir) / "ocr_for_tables.pdf")
    subprocess.run(["ocrmypdf", "--skip-text", "-l", "fra+eng", "--output-type", "pdf", "--optimize", "0", src, out],
                   check=True, capture_output=True, timeout=240)
    return out


def _text_fallback_rows(pdf) -> list[list[str]]:
    rows: list[list[str]] = []
    for pno, page in enumerate(pdf.pages, start=1):
        text = page.extract_text(layout=True) or ""
        lines = [ln for ln in text.splitlines() if ln.strip()]
        if not lines:
            continue
        if len(pdf.pages) > 1:
            rows.append([f"— Page {pno} —"])
        rows.extend([c for c in re.split(r"\s{2,}", ln.strip())] for ln in lines)
    return rows


def _write_sheet(ws, rows: list[list[str]]) -> None:
    from openpyxl.styles import Font, PatternFill

    cells = [c for r in rows for c in r]
    decimal = detect_decimal(cells)
    header = bool(rows) and len(rows) > 1 and is_header(rows[0], decimal)
    widths: dict[int, int] = {}
    for i, row in enumerate(rows, start=1):
        for j, raw in enumerate(row, start=1):
            cell = ws.cell(row=i, column=j)
            typed = None if (header and i == 1) else type_cell(raw, decimal)
            if typed:
                cell.value, cell.number_format = typed.value, typed.number_format
            else:
                cell.value = raw or None
            widths[j] = max(widths.get(j, 0), len(raw))
    for j, w in widths.items():
        ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = min(60, max(8, w + 2))
    if header:
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="EDE9FE")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions


def convert(src: str, out: str, tmp_dir: str) -> None:
    import pdfplumber
    from openpyxl import Workbook

    src = _ocr_if_scanned(src, tmp_dir)
    wb = Workbook()
    wb.remove(wb.active)
    with pdfplumber.open(src) as pdf:
        tables: list[Table] = []
        for pno, page in enumerate(pdf.pages, start=1):
            tables.extend(_page_tables(page, pno))
            page.flush_cache()  # keep memory flat on long documents
        tables = merge_continued(tables)
        for i, t in enumerate(tables, start=1):
            pages = f"p{t.page}" if t.page == t.last_page else f"p{t.page}-{t.last_page}"
            _write_sheet(wb.create_sheet(f"Table {i} ({pages})"[:31]), t.rows)
        if not tables:
            _write_sheet(wb.create_sheet("Text"), _text_fallback_rows(pdf))
    wb.save(out)
