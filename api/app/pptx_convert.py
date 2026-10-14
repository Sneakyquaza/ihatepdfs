"""PDF -> editable PowerPoint.

Each slide = the page artwork (images, vector graphics) as a background picture, with every text
line re-created on top as a real text box (font, size, color, bold/italic, angle).

Coordinates: PyMuPDF reports text boxes and directions in *unrotated* page space; the slide shows
the page as displayed, so everything goes through `page.rotation_matrix` (see `line_frame`).
"""

import math
from dataclasses import dataclass
from pathlib import Path

import pymupdf

EMU_PER_PT = 12700
_FONT_ALIASES = {"helvetica": "Arial", "arialmt": "Arial", "times": "Times New Roman", "timesnewromanpsmt": "Times New Roman",
                 "courier": "Courier New", "couriernewpsmt": "Courier New", "symbol": "Symbol"}


def pptx_font(name: str) -> str:
    """'ABCDEF+Calibri-Bold' -> 'Calibri'; PDF base-14 names -> common Office fonts."""
    name = name.split("+", 1)[-1]
    family = name.split("-", 1)[0].split(",", 1)[0]
    alias = _FONT_ALIASES.get(family.lower())
    if alias:
        return alias
    for suffix in ("PSMT", "MT", "PS"):
        if family.endswith(suffix) and len(family) > len(suffix):
            family = family[: -len(suffix)]
            break
    return family or "Arial"


def xml_safe(text: str) -> str:
    return "".join(ch for ch in text if ch == "\t" or ord(ch) >= 0x20)


@dataclass(frozen=True)
class Frame:
    """A text line's box in displayed page space, before rotation, plus its clockwise angle."""
    cx: float
    cy: float
    length: float  # along the text direction
    thickness: float  # across it
    angle: float  # degrees, clockwise (PowerPoint's convention), in [0, 360)


def line_frame(bbox: tuple, direction: tuple, font_size: float, rotation_matrix: pymupdf.Matrix) -> Frame:
    """Unrotated bbox + writing direction -> centred, un-rotated frame and angle as seen on the displayed page."""
    rect = pymupdf.Rect(bbox) * rotation_matrix
    origin = pymupdf.Point(0, 0) * rotation_matrix
    d = pymupdf.Point(direction) * rotation_matrix - origin
    angle = math.degrees(math.atan2(d.y, d.x)) % 360  # y points down -> clockwise
    if min(angle, 360 - angle) < 0.5:
        angle = 0.0
    c, s = abs(math.cos(math.radians(angle))), abs(math.sin(math.radians(angle)))
    # the line's bbox is the axis-aligned hull of a length x thickness box rotated by `angle`
    thickness = min(font_size * 1.2, rect.width * s + rect.height * c) or font_size
    if c >= s:
        length = (rect.width - thickness * s) / max(c, 1e-6)
    else:
        length = (rect.height - thickness * c) / max(s, 1e-6)
    return Frame(cx=(rect.x0 + rect.x1) / 2, cy=(rect.y0 + rect.y1) / 2, length=max(length, font_size / 2),
                 thickness=thickness, angle=angle)


def _text_lines(page: pymupdf.Page) -> list[tuple[Frame, list[dict]]]:
    lines = []
    for block in page.get_text("dict", flags=pymupdf.TEXTFLAGS_TEXT)["blocks"]:
        for line in block.get("lines", []):
            spans = [sp for sp in line["spans"] if sp["text"].strip()]
            if not spans:
                continue
            size = max(sp["size"] for sp in spans)
            lines.append((line_frame(line["bbox"], line["dir"], size, page.rotation_matrix), line["spans"]))
    return lines


def _strip_text(page: pymupdf.Page) -> None:
    """Remove all text from the page, keeping images and vector art for the background picture."""
    for block in page.get_text("dict", flags=pymupdf.TEXTFLAGS_TEXT)["blocks"]:
        for line in block.get("lines", []):
            for sp in line["spans"]:
                page.add_redact_annot(pymupdf.Rect(sp["bbox"]), fill=False)
    page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE, graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
                          text=pymupdf.PDF_REDACT_TEXT_REMOVE)


def convert(src: str, out: str, tmp_dir: str) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.util import Emu, Pt

    def emu(pt: float) -> Emu:
        return Emu(int(pt * EMU_PER_PT))

    prs = Presentation()
    with pymupdf.open(src) as doc:
        # PowerPoint has a single slide size: use the first page's (as displayed) and letterbox the others
        first = doc[0].rect
        prs.slide_width, prs.slide_height = emu(first.width), emu(first.height)
        blank = prs.slide_layouts[6]
        for i, page in enumerate(doc):
            # uniform scale + centring keeps proportions for pages of another size/orientation
            scale = min(first.width / page.rect.width, first.height / page.rect.height)
            ox = (first.width - page.rect.width * scale) / 2
            oy = (first.height - page.rect.height * scale) / 2
            lines = _text_lines(page)
            if lines:
                _strip_text(page)

            slide = prs.slides.add_slide(blank)
            img = Path(tmp_dir) / f"slide{i}.png"
            page.get_pixmap(dpi=150).save(img)
            slide.shapes.add_picture(str(img), emu(ox), emu(oy), emu(page.rect.width * scale), emu(page.rect.height * scale))

            for f, spans in lines:
                # a little slack so the line never wraps; kept symmetric so rotation stays centred
                length = f.length * 1.08 + 4
                box = slide.shapes.add_textbox(
                    emu(ox + (f.cx - length / 2) * scale), emu(oy + (f.cy - f.thickness / 2) * scale),
                    emu(length * scale), emu(f.thickness * scale),
                )
                box.rotation = f.angle
                tf = box.text_frame
                tf.word_wrap = False
                tf.auto_size = MSO_AUTO_SIZE.NONE
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                para = tf.paragraphs[0]
                for sp in spans:
                    text = xml_safe(sp["text"])
                    if not text:
                        continue
                    run = para.add_run()
                    run.text = text
                    font = run.font
                    font.size = Pt(max(1.0, sp["size"] * scale))
                    font.bold = bool(sp["flags"] & 16)
                    font.italic = bool(sp["flags"] & 2)
                    font.name = pptx_font(sp["font"])
                    c = sp["color"]
                    font.color.rgb = RGBColor((c >> 16) & 255, (c >> 8) & 255, c & 255)
    prs.save(out)
