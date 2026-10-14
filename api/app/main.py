"""Server-side processing for the heavy tools (conversions, compression, OCR, security, compare, upscale)."""

import asyncio
import difflib
import os
import re
from pathlib import Path
from typing import Awaitable, Callable

import fitz  # PyMuPDF
import pikepdf
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response

from .ratelimit import rate_limit_middleware
from .redact import PRESETS, parse_boxes, redact
from .util import Workspace, assert_public_url, run, save_upload, send, stem_of

app = FastAPI(title="iHatePDFs API", docs_url=None, redoc_url=None, openapi_url=None)
app.middleware("http")(rate_limit_middleware)
app.add_middleware(
    CORSMiddleware,
    allow_origins=os.getenv("CORS_ORIGINS", "http://localhost:3000").split(","),
    allow_methods=["*"],
    allow_headers=["*"],
    expose_headers=["Content-Disposition"],
)

PDF = {".pdf"}
OFFICE = {".doc", ".docx", ".odt", ".rtf", ".txt", ".ppt", ".pptx", ".odp", ".xls", ".xlsx", ".ods", ".csv"}
IMAGES = {".png", ".jpg", ".jpeg", ".webp", ".bmp"}
OCR_LANGS = {"eng", "fra", "deu", "spa", "ita", "por", "nld"}

# Limit concurrent heavy jobs so one burst can't exhaust CPU/RAM, and cap the queue behind them.
HEAVY = asyncio.Semaphore(int(os.getenv("MAX_HEAVY_JOBS", "2")))
MAX_QUEUE = int(os.getenv("MAX_QUEUE", "20"))
_queued = 0


async def handle(work: Callable[[Workspace], Awaitable[Response]]) -> Response:
    global _queued
    if _queued >= MAX_QUEUE:
        raise HTTPException(503, "Server busy, please retry in a minute.", headers={"Retry-After": "30"})
    ws = Workspace()
    _queued += 1
    waiting = True
    try:
        async with HEAVY:
            _queued -= 1
            waiting = False
            return await work(ws)
    except BaseException:
        ws.cleanup()
        raise
    finally:
        if waiting:  # cancelled (client gone) before getting a slot
            _queued -= 1


@app.get("/api/health")
async def health():
    return {"ok": True}


# ---------- Compression / repair ----------

GS_LEVELS = {"low": "/printer", "recommended": "/ebook", "extreme": "/screen"}


@app.post("/api/compress-pdf")
async def compress_pdf(file: UploadFile = File(...), level: str = Form("recommended")):
    if level not in GS_LEVELS:
        raise HTTPException(400, "Invalid level")

    async def work(ws: Workspace):
        src = await save_upload(file, ws, PDF)
        out = ws.file(".pdf")
        await run([
            "gs", "-sDEVICE=pdfwrite", "-dCompatibilityLevel=1.5", f"-dPDFSETTINGS={GS_LEVELS[level]}",
            "-dNOPAUSE", "-dQUIET", "-dBATCH", "-dSAFER", "-dDetectDuplicateImages=true",
            f"-sOutputFile={out}", str(src),
        ])
        best = out if out.stat().st_size < src.stat().st_size else src
        return send(best, f"{stem_of(file)}_compressed.pdf", ws)

    return await handle(work)


@app.post("/api/repair-pdf")
async def repair_pdf(file: UploadFile = File(...)):
    async def work(ws: Workspace):
        src = await save_upload(file, ws, PDF)
        out = ws.file(".pdf")
        try:
            def _repair():
                with pikepdf.open(src) as pdf:
                    pdf.save(out)
            await asyncio.to_thread(_repair)
        except Exception:
            await run(["gs", "-sDEVICE=pdfwrite", "-dNOPAUSE", "-dQUIET", "-dBATCH", "-dSAFER",
                       f"-sOutputFile={out}", str(src)])
        return send(out, f"{stem_of(file)}_repaired.pdf", ws)

    return await handle(work)


# ---------- Conversions ----------

@app.post("/api/office-to-pdf")
async def office_to_pdf(file: UploadFile = File(...)):
    async def work(ws: Workspace):
        src = await save_upload(file, ws, OFFICE)
        profile = ws.path / "lo_profile"
        await run([
            "soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--norestore",
            "--convert-to", "pdf", "--outdir", str(ws.path), str(src),
        ], timeout=240, env={"HOME": str(ws.path)})
        return send(src.with_suffix(".pdf"), f"{stem_of(file)}.pdf", ws)

    return await handle(work)


@app.post("/api/pdf-to-word")
async def pdf_to_word(file: UploadFile = File(...)):
    async def work(ws: Workspace):
        from pdf2docx import Converter

        src = await save_upload(file, ws, PDF)
        out = ws.file(".docx")

        def _convert():
            cv = Converter(str(src))
            try:
                cv.convert(str(out))
            finally:
                cv.close()

        await asyncio.to_thread(_convert)
        return send(out, f"{stem_of(file)}.docx", ws)

    return await handle(work)


@app.post("/api/pdf-to-excel")
async def pdf_to_excel(file: UploadFile = File(...)):
    async def work(ws: Workspace):
        import pdfplumber
        from openpyxl import Workbook

        src = await save_upload(file, ws, PDF)
        out = ws.file(".xlsx")

        def _convert():
            wb = Workbook()
            wb.remove(wb.active)
            with pdfplumber.open(src) as pdf:
                for pno, page in enumerate(pdf.pages, start=1):
                    tables = page.extract_tables()
                    if tables:
                        for tno, table in enumerate(tables, start=1):
                            ws_ = wb.create_sheet(f"P{pno}-T{tno}"[:31])
                            for row in table:
                                ws_.append([c if c is not None else "" for c in row])
                    else:
                        text = page.extract_text() or ""
                        if text.strip():
                            ws_ = wb.create_sheet(f"P{pno}-text"[:31])
                            for line in text.splitlines():
                                ws_.append(line.split())
            if not wb.sheetnames:
                wb.create_sheet("Empty")
            wb.save(out)

        await asyncio.to_thread(_convert)
        return send(out, f"{stem_of(file)}.xlsx", ws)

    return await handle(work)


@app.post("/api/pdf-to-powerpoint")
async def pdf_to_powerpoint(file: UploadFile = File(...)):
    async def work(ws: Workspace):
        from pptx import Presentation
        from pptx.util import Emu

        src = await save_upload(file, ws, PDF)
        out = ws.file(".pptx")

        def _convert():
            prs = Presentation()
            with fitz.open(src) as doc:
                if doc.page_count == 0:
                    raise HTTPException(400, "PDF has no pages")
                first = doc[0].rect
                # 1pt = 12700 EMU
                prs.slide_width, prs.slide_height = Emu(int(first.width * 12700)), Emu(int(first.height * 12700))
                blank = prs.slide_layouts[6]
                for page in doc:
                    img = ws.file(".png")
                    page.get_pixmap(dpi=150).save(img)
                    slide = prs.slides.add_slide(blank)
                    slide.shapes.add_picture(str(img), 0, 0, prs.slide_width, prs.slide_height)
            prs.save(out)

        await asyncio.to_thread(_convert)
        return send(out, f"{stem_of(file)}.pptx", ws)

    return await handle(work)


def _safe_fetcher():
    """WeasyPrint fetcher limited to data: and public http(s) URLs, re-checking every redirect hop."""
    from urllib import request

    from weasyprint.urls import URLFetcher

    class _CheckedRedirects(request.HTTPRedirectHandler):
        def redirect_request(self, req, fp, code, msg, headers, newurl):
            try:
                assert_public_url(newurl)
            except HTTPException as e:
                raise ValueError(e.detail)
            return super().redirect_request(req, fp, code, msg, headers, newurl)

    class SafeFetcher(URLFetcher):
        def __init__(self):
            super().__init__(timeout=15, allowed_protocols={"http", "https", "data"}, allow_redirects=False)
            self.add_handler(_CheckedRedirects())

        def fetch(self, url, headers=None):
            if not url.startswith("data:"):
                try:
                    assert_public_url(url)
                except HTTPException as e:
                    raise ValueError(e.detail)
            return super().fetch(url, headers)

    return SafeFetcher()


@app.post("/api/html-to-pdf")
async def html_to_pdf(file: UploadFile | None = File(None), url: str | None = Form(None)):
    if not file and not url:
        raise HTTPException(400, "Provide an HTML file or a URL")

    async def work(ws: Workspace):
        from weasyprint import HTML

        out = ws.file(".pdf")
        if url:
            assert_public_url(url)
            await asyncio.to_thread(lambda: HTML(url=url, url_fetcher=_safe_fetcher()).write_pdf(out))
            name = "webpage.pdf"
        else:
            src = await save_upload(file, ws, {".html", ".htm"})
            html = src.read_text(errors="replace")
            await asyncio.to_thread(lambda: HTML(string=html, url_fetcher=_safe_fetcher()).write_pdf(out))
            name = f"{stem_of(file)}.pdf"
        return send(out, name, ws)

    return await handle(work)


# ---------- OCR / PDF/A ----------

@app.post("/api/ocr-pdf")
async def ocr_pdf(file: UploadFile = File(...), lang: str = Form("eng")):
    langs = [l for l in lang.split("+") if l]
    if not langs or any(l not in OCR_LANGS for l in langs):
        raise HTTPException(400, f"Invalid language. Allowed: {', '.join(sorted(OCR_LANGS))}")

    async def work(ws: Workspace):
        src = await save_upload(file, ws, PDF)
        out = ws.file(".pdf")
        await run(["ocrmypdf", "--skip-text", "--rotate-pages", "--deskew", "-l", "+".join(langs),
                   "--output-type", "pdf", str(src), str(out)], timeout=600)
        return send(out, f"{stem_of(file)}_ocr.pdf", ws)

    return await handle(work)


@app.post("/api/pdf-to-pdfa")
async def pdf_to_pdfa(file: UploadFile = File(...)):
    async def work(ws: Workspace):
        src = await save_upload(file, ws, PDF)
        out = ws.file(".pdf")
        # --tesseract-timeout 0 disables OCR while still producing a PDF/A
        await run(["ocrmypdf", "--skip-text", "--tesseract-timeout", "0", "--output-type", "pdfa-2",
                   str(src), str(out)], timeout=300)
        return send(out, f"{stem_of(file)}_pdfa.pdf", ws)

    return await handle(work)


# ---------- Security ----------

@app.post("/api/protect-pdf")
async def protect_pdf(file: UploadFile = File(...), password: str = Form(...)):
    if not password:
        raise HTTPException(400, "Password required")

    async def work(ws: Workspace):
        src = await save_upload(file, ws, PDF)
        out = ws.file(".pdf")

        def _protect():
            with pikepdf.open(src) as pdf:
                pdf.save(out, encryption=pikepdf.Encryption(owner=password, user=password, R=6))

        await asyncio.to_thread(_protect)
        return send(out, f"{stem_of(file)}_protected.pdf", ws)

    return await handle(work)


@app.post("/api/unlock-pdf")
async def unlock_pdf(file: UploadFile = File(...), password: str = Form("")):
    async def work(ws: Workspace):
        src = await save_upload(file, ws, PDF)
        out = ws.file(".pdf")

        def _unlock():
            try:
                with pikepdf.open(src, password=password) as pdf:
                    pdf.save(out)
            except pikepdf.PasswordError:
                raise HTTPException(400, "Wrong password")

        await asyncio.to_thread(_unlock)
        return send(out, f"{stem_of(file)}_unlocked.pdf", ws)

    return await handle(work)


# ---------- Compare ----------

def _words(doc: fitz.Document) -> list[tuple[int, fitz.Rect, str]]:
    return [(pno, fitz.Rect(w[:4]), w[4]) for pno, page in enumerate(doc) for w in page.get_text("words")]


@app.post("/api/compare-pdf")
async def compare_pdf(file_a: UploadFile = File(...), file_b: UploadFile = File(...)):
    async def work(ws: Workspace):
        src_a = await save_upload(file_a, ws, PDF)
        src_b = await save_upload(file_b, ws, PDF)
        out = ws.file(".pdf")

        def _compare():
            with fitz.open(src_a) as a, fitz.open(src_b) as b, fitz.open() as result:
                wa, wb = _words(a), _words(b)
                removed: dict[int, list[fitz.Rect]] = {}
                added: dict[int, list[fitz.Rect]] = {}
                sm = difflib.SequenceMatcher(None, [w[2] for w in wa], [w[2] for w in wb], autojunk=False)
                for tag, i1, i2, j1, j2 in sm.get_opcodes():
                    if tag in ("delete", "replace"):
                        for pno, rect, _ in wa[i1:i2]:
                            removed.setdefault(pno, []).append(rect)
                    if tag in ("insert", "replace"):
                        for pno, rect, _ in wb[j1:j2]:
                            added.setdefault(pno, []).append(rect)

                gap, header = 24, 28
                for i in range(max(a.page_count, b.page_count)):
                    ra = a[i].rect if i < a.page_count else fitz.Rect(0, 0, 0, 0)
                    rb = b[i].rect if i < b.page_count else fitz.Rect(0, 0, 0, 0)
                    width = ra.width + rb.width + gap
                    height = max(ra.height, rb.height) + header
                    page = result.new_page(width=width, height=height)
                    page.insert_text((8, 18), f"A: {file_a.filename}  (page {i + 1})", fontsize=10, color=(0.7, 0, 0))
                    page.insert_text((ra.width + gap + 8, 18), f"B: {file_b.filename}  (page {i + 1})",
                                     fontsize=10, color=(0, 0.5, 0))
                    if i < a.page_count:
                        page.show_pdf_page(fitz.Rect(0, header, ra.width, header + ra.height), a, i)
                        for r in removed.get(i, []):
                            page.draw_rect(r + (0, header, 0, header), color=None, fill=(1, 0, 0), fill_opacity=0.3)
                    if i < b.page_count:
                        ox = ra.width + gap
                        page.show_pdf_page(fitz.Rect(ox, header, ox + rb.width, header + rb.height), b, i)
                        for r in added.get(i, []):
                            page.draw_rect(r + (ox, header, ox, header), color=None, fill=(0, 0.8, 0), fill_opacity=0.3)
                    page.draw_line((ra.width + gap / 2, 0), (ra.width + gap / 2, height), color=(0.6, 0.6, 0.6))
                result.save(out)

        await asyncio.to_thread(_compare)
        return send(out, "comparison.pdf", ws)

    return await handle(work)


# ---------- Images ----------

UPSCALE_MAX_INPUT_PIXELS = {2: 2_500_000, 4: 800_000}


@app.post("/api/upscale-image")
async def upscale_image(file: UploadFile = File(...), scale: int = Form(2)):
    if scale not in UPSCALE_MAX_INPUT_PIXELS:
        raise HTTPException(400, "Scale must be 2 or 4")

    async def work(ws: Workspace):
        from PIL import Image, ImageFilter

        src = await save_upload(file, ws, IMAGES)
        out = ws.file(".png")

        def _upscale():
            with Image.open(src) as im:
                if im.width * im.height > UPSCALE_MAX_INPUT_PIXELS[scale]:
                    raise HTTPException(400, f"Image too large for x{scale} upscale")
                im = im.convert("RGBA") if im.mode in ("RGBA", "LA", "P") else im.convert("RGB")
                model = Path(os.getenv("MODELS_DIR", "/models")) / f"FSRCNN_x{scale}.pb"
                try:
                    import cv2
                    import numpy as np

                    if not model.exists():
                        raise FileNotFoundError(model)
                    sr = cv2.dnn_superres.DnnSuperResImpl_create()
                    sr.readModel(str(model))
                    sr.setModel("fsrcnn", scale)
                    rgb = np.array(im.convert("RGB"))[:, :, ::-1]
                    up = Image.fromarray(sr.upsample(rgb)[:, :, ::-1])
                    if im.mode == "RGBA":
                        up.putalpha(im.getchannel("A").resize(up.size, Image.LANCZOS))
                except (ImportError, AttributeError, FileNotFoundError):
                    up = im.resize((im.width * scale, im.height * scale), Image.LANCZOS)
                    up = up.filter(ImageFilter.UnsharpMask(radius=2, percent=80, threshold=2))
                up.save(out, optimize=True)

        await asyncio.to_thread(_upscale)
        return send(out, f"{stem_of(file)}_x{scale}.png", ws)

    return await handle(work)


# ---------- Redaction / cleanup ----------

@app.post("/api/redact-pdf")
async def redact_pdf(
    file: UploadFile = File(...),
    editor: str | None = Form(None),
    terms: str = Form(""),
    redactEmails: bool = Form(False),
    redactPhones: bool = Form(False),
    redactIban: bool = Form(False),
    redactCards: bool = Form(False),
):
    boxes = parse_boxes(editor)
    term_list = [t.strip() for t in re.split(r"[,\n]", terms) if t.strip()][:200]
    flags = {"emails": redactEmails, "phones": redactPhones, "iban": redactIban, "cards": redactCards}
    presets = [p for p, on in flags.items() if on and p in PRESETS]

    async def work(ws: Workspace):
        src = await save_upload(file, ws, PDF)
        out = ws.file(".pdf")
        count = await asyncio.to_thread(redact, str(src), str(out), boxes, term_list, presets)
        resp = send(out, f"{stem_of(file)}_redacted.pdf", ws)
        resp.headers["X-Redaction-Count"] = str(count)
        return resp

    return await handle(work)


@app.post("/api/flatten-pdf")
async def flatten_pdf(file: UploadFile = File(...)):
    async def work(ws: Workspace):
        src = await save_upload(file, ws, PDF)
        out = ws.file(".pdf")

        def _flatten():
            with pikepdf.open(src) as pdf:
                # form fields need appearance streams before they can be burnt into the page
                pdf.generate_appearance_streams()
                pdf.flatten_annotations("all")
                if "/AcroForm" in pdf.Root:
                    del pdf.Root.AcroForm
                pdf.save(out)

        await asyncio.to_thread(_flatten)
        return send(out, f"{stem_of(file)}_flat.pdf", ws)

    return await handle(work)


@app.post("/api/grayscale-pdf")
async def grayscale_pdf(file: UploadFile = File(...)):
    async def work(ws: Workspace):
        src = await save_upload(file, ws, PDF)
        out = ws.file(".pdf")
        await run([
            "gs", "-sDEVICE=pdfwrite", "-sColorConversionStrategy=Gray", "-dProcessColorModel=/DeviceGray",
            "-dOverrideICC", "-dNOPAUSE", "-dQUIET", "-dBATCH", "-dSAFER", f"-sOutputFile={out}", str(src),
        ])
        return send(out, f"{stem_of(file)}_gray.pdf", ws)

    return await handle(work)


@app.post("/api/image-to-text")
async def image_to_text(file: UploadFile = File(...), lang: str = Form("eng")):
    langs = [l for l in lang.split("+") if l]
    if not langs or any(l not in OCR_LANGS for l in langs):
        raise HTTPException(400, f"Invalid language. Allowed: {', '.join(sorted(OCR_LANGS))}")

    async def work(ws: Workspace):
        src = await save_upload(file, ws, IMAGES)
        base = ws.file("")
        await run(["tesseract", str(src), str(base), "-l", "+".join(langs)], timeout=180)
        return send(base.with_suffix(".txt"), f"{stem_of(file)}.txt", ws)

    return await handle(work)

